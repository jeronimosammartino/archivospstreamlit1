import io
from datetime import datetime
import streamlit as st
import pandas as pd
import openpyxl
import plotly.express as px
import plotly.graph_objects as go
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

st.set_page_config(page_title="Control de Métodos, Costos y HH", layout="wide")

st.title("📊 Control de Métodos: Horas, Costos y Rendimientos")
st.write("Tablero integral de control con exportación completa a PowerPoint.")

uploaded_file = st.file_uploader("Sube la matriz del proyecto (.xlsx)", type=["xlsx", "xls"])

def leer_fabricost(file_bytes):
    try:
        df_costos = pd.read_excel(file_bytes, sheet_name="FABRICOST")
        col_recurso, col_tarifa = None, None
        for col in df_costos.columns:
            col_str = str(col).strip().upper()
            if "RECURSO" in col_str:
                col_recurso = col
            elif any(k in col_str for k in ["COSTO", "CIF", "TARIFA", "VALOR"]):
                col_tarifa = col
        if col_recurso and col_tarifa:
            df_costos = df_costos[[col_recurso, col_tarifa]].dropna()
            df_costos.columns = ["Recurso", "Tarifa"]
            df_costos["Recurso_Clean"] = df_costos["Recurso"].astype(str).str.strip().str.upper()
            df_costos["Tarifa"] = pd.to_numeric(df_costos["Tarifa"], errors="coerce").fillna(0)
            return df_costos
    except Exception:
        pass
    return pd.DataFrame(columns=["Recurso_Clean", "Tarifa"])

def leer_matriz_exacta(file_bytes, sheet_name):
    wb = openpyxl.load_workbook(file_bytes, data_only=True)
    ws = wb[sheet_name]
    
    for mr in list(ws.merged_cells.ranges):
        val = ws.cell(row=mr.min_row, column=mr.min_col).value
        ws.unmerge_cells(range_string=str(mr))
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                ws.cell(row=r, column=c).value = val
                
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return pd.DataFrame()
        
    header_idx = 0
    for idx, row in enumerate(filas[:10]):
        textos = [str(c).strip().upper() for c in row if c is not None]
        if any("CATEGORIA" in t or "GDF" in t or "PESO" in t for t in textos):
            header_idx = idx
            break
            
    headers = [str(c).strip().replace("\n", " ") if c is not None else f"Col_{i}" for i, c in enumerate(filas[header_idx])]
    df = pd.DataFrame(filas[header_idx + 1:], columns=headers).dropna(how="all").reset_index(drop=True)
    
    if "Categoria" in df.columns:
        df["Categoria"] = df["Categoria"].ffill()
    return df

def procesar_datos(df, df_tarifas):
    cols_hh = [c for c in df.columns if c.startswith("HH ") or c in ["Autycontrol", "MET"]]
    df_det = df[~df['GDF'].astype(str).str.contains("Total", case=False, na=False)].copy()
    
    def normalizar_cat(val):
        v = str(val).upper()
        if "OFERTA" in v:
            return "OFERTADAS"
        elif "HDR" in v:
            return "ESTIMADAS"
        elif "REAL" in v:
            return "REALES"
        return None

    df_det['Estado'] = df_det['Categoria'].apply(normalizar_cat)
    df_det = df_det.dropna(subset=['Estado'])
    
    df_long = df_det.melt(
        id_vars=['Estado', 'GDF', 'Peso'],
        value_vars=cols_hh,
        var_name='Recurso',
        value_name='Horas'
    )
    df_long['Horas'] = pd.to_numeric(df_long['Horas'], errors='coerce').fillna(0)
    df_long['Peso'] = pd.to_numeric(df_long['Peso'], errors='coerce').fillna(0)
    
    df_long['Recurso_Clean'] = df_long['Recurso'].astype(str).str.strip().str.upper()
    if not df_tarifas.empty:
        df_long = pd.merge(df_long, df_tarifas[['Recurso_Clean', 'Tarifa']], on='Recurso_Clean', how='left')
        df_long['Tarifa'] = df_long['Tarifa'].fillna(0)
        df_long['Costo'] = df_long['Horas'] * df_long['Tarifa']
    else:
        df_long['Costo'] = 0.0
        
    df_pesos = df_det[['Estado', 'GDF', 'Peso']].drop_duplicates()
    df_pesos['Peso'] = pd.to_numeric(df_pesos['Peso'], errors='coerce').fillna(0)
    
    return df_long, df_pesos

def formato_m_k(val):
    if abs(val) >= 1_000_000:
        return f"{val/1_000_000:.2f}M"
    elif abs(val) >= 1_000:
        return f"{val/1_000:.2f}K"
    return f"{val:.2f}"

def setup_ax():
    fig, ax = plt.subplots(figsize=(8.2, 4.8), dpi=160)
    fig.patch.set_facecolor('#FFFFFF')
    ax.set_facecolor('#F8FAFC')
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['left'].set_color('#CBD5E1')
    ax.spines['bottom'].set_color('#CBD5E1')
    ax.grid(axis='y', linestyle='--', alpha=0.6, color='#E2E8F0')
    ax.set_axisbelow(True)
    return fig, ax

def export_fig(fig):
    plt.tight_layout()
    buf = io.BytesIO()
    plt.savefig(buf, format='png', dpi=160)
    plt.close(fig)
    buf.seek(0)
    return buf

def plot_barra_global(estados, valores, es_costo=False, ylabel=""):
    fig, ax = setup_ax()
    colores = ["#F59E0B", "#06B6D4", "#6366F1"]
    bars = ax.bar(estados, valores, color=colores, width=0.52)
    max_v = max(valores) if max(valores) > 0 else 1
    for b in bars:
        h = b.get_height()
        label = formato_m_k(h) if es_costo else (f"{h:.2f}" if "Kg" in ylabel else f"{int(h):,}")
        ax.text(b.get_x() + b.get_width()/2., h + (max_v * 0.02), label, ha='center', va='bottom', fontsize=9.5, fontweight='bold', color='#1E293B')
    ax.set_ylabel(ylabel, fontsize=9.5, color='#475569')
    return export_fig(fig)

def plot_barras_agrupadas(df_base, col_x, metrica="Horas", es_costo=False):
    fig, ax = setup_ax()
    df_piv = df_base.pivot_table(index=col_x, columns='Estado', values=metrica, aggfunc='sum').fillna(0)
    estados_ord = [e for e in ["OFERTADAS", "ESTIMADAS", "REALES"] if e in df_piv.columns]
    df_piv = df_piv[estados_ord]
    
    x = np.arange(len(df_piv))
    width = 0.25
    colores = {"OFERTADAS": "#F59E0B", "ESTIMADAS": "#06B6D4", "REALES": "#6366F1"}
    
    for i, est in enumerate(estados_ord):
        ax.bar(x + (i - 1) * width, df_piv[est], width, label=est, color=colores.get(est, "#94A3B8"))
        
    ax.set_xticks(x)
    labels = [str(lbl)[:16] for lbl in df_piv.index]
    ax.set_xticklabels(labels, rotation=40, ha='right', fontsize=8, color='#334155')
    ax.legend(frameon=True, facecolor='#FFFFFF', edgecolor='#E2E8F0', fontsize=8.5)
    ax.set_ylabel("Costo ($)" if es_costo else "Horas Hombre (HH)", fontsize=9, color='#475569')
    return export_fig(fig)

def plot_cascada(df_base, col_agrupacion, metrica, estado_fin, estado_ini):
    fig, ax = setup_ax()
    piv = df_base.pivot_table(index=col_agrupacion, columns='Estado', values=metrica, aggfunc='sum').fillna(0)
    if estado_fin in piv.columns and estado_ini in piv.columns:
        piv['Desvio'] = piv[estado_fin] - piv[estado_ini]
        piv = piv.sort_values(by='Desvio', ascending=False)
        
        nombres = [str(n)[:15] for n in piv.index.tolist()] + ["Total"]
        deltas = piv['Desvio'].tolist() + [piv['Desvio'].sum()]
        
        bottoms = []
        curr = 0
        for i, val in enumerate(deltas[:-1]):
            if val >= 0:
                bottoms.append(curr)
                curr += val
            else:
                curr += val
                bottoms.append(curr)
        bottoms.append(0)
        
        bar_colors = ["#EF4444" if val > 0 else "#10B981" for val in deltas[:-1]] + ["#06B6D4"]
        
        x = np.arange(len(nombres))
        ax.bar(x, [abs(v) if i < len(deltas)-1 else v for i, v in enumerate(deltas)], bottom=bottoms, color=bar_colors, width=0.55)
        ax.set_xticks(x)
        ax.set_xticklabels(nombres, rotation=45, ha='right', fontsize=7.5, color='#334155')
        ax.set_ylabel("Variación " + ("($)" if metrica == "Costo" else "(HH)"), fontsize=9, color='#475569')
    return export_fig(fig)

def generar_powerpoint_completo(df_f, metrics, hoja_nombre):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    # 1. PORTADA EJECUTIVA
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(15, 23, 42)
    bg1.line.fill.background()
    
    tb = slide1.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(11), Inches(2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "INFORME EJECUTIVO DE CONTROL OPERACIONAL"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = tf.add_paragraph()
    p2.text = f"Balance de Horas Hombre (HH), Costos y Rendimientos | Proyecto: {hoja_nombre}"
    p2.font.size = Pt(17)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    
    tb_meta = slide1.shapes.add_textbox(Inches(1.2), Inches(5.6), Inches(10), Inches(1))
    tf_m = tb_meta.text_frame
    pm = tf_m.paragraphs[0]
    pm.text = f"Fecha de emisión: {datetime.now().strftime('%d/%m/%Y')}   |   Gerencia de Métodos y Procesos   |   Confidencial"
    pm.font.size = Pt(11)
    pm.font.color.rgb = RGBColor(100, 116, 139)
    
    def agregar_diapositiva_grafica(titulo, subtitulo, chart_buf, seccion_badge):
        slide = prs.slides.add_slide(blank_layout)
        
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(248, 250, 252)
        bg.line.fill.background()
        
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.15))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = RGBColor(15, 23, 42)
        top_bar.line.fill.background()
        
        tb_t = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.5), Inches(0.9))
        tf_t = tb_t.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = titulo
        p_t.font.size = Pt(19)
        p_t.font.bold = True
        p_t.font.color.rgb = RGBColor(255, 255, 255)
        
        p_sub = tf_t.add_paragraph()
        p_sub.text = f"{seccion_badge}  |  {subtitulo}"
        p_sub.font.size = Pt(10.5)
        p_sub.font.color.rgb = RGBColor(148, 163, 184)
        
        slide.shapes.add_picture(chart_buf, Inches(0.6), Inches(1.5), width=Inches(8.4))
        
        box_n = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.2), Inches(1.5), Inches(3.5), Inches(5.3))
        box_n.fill.solid()
        box_n.fill.fore_color.rgb = RGBColor(255, 255, 255)
        box_n.line.color.rgb = RGBColor(203, 213, 225)
        
        tb_n = slide.shapes.add_textbox(Inches(9.4), Inches(1.7), Inches(3.1), Inches(4.9))
        tf_n = tb_n.text_frame
        tf_n.word_wrap = True
        
        p_nh = tf_n.paragraphs[0]
        p_nh.text = "NOTAS Y CONCLUSIONES"
        p_nh.font.size = Pt(12)
        p_nh.font.bold = True
        p_nh.font.color.rgb = RGBColor(30, 41, 59)
        
        p_sp = tf_n.add_paragraph()
        p_sp.text = "──────────────────────"
        p_sp.font.size = Pt(9)
        p_sp.font.color.rgb = RGBColor(203, 213, 225)
        
        p_nt = tf_n.add_paragraph()
        p_nt.text = "• Proyecto: " + str(hoja_nombre)
        p_nt.font.size = Pt(10)
        p_nt.font.color.rgb = RGBColor(71, 85, 105)
        
        p_nt2 = tf_n.add_paragraph()
        p_nt2.text = "• Ingrese aquí notas técnicas o puntos clave de discusión para esta gráfica:"
        p_nt2.font.size = Pt(9.5)
        p_nt2.font.color.rgb = RGBColor(100, 116, 139)
        
        p_edit = tf_n.add_paragraph()
        p_edit.text = "\n[Texto editable...]"
        p_edit.font.size = Pt(9.5)
        p_edit.font.color.rgb = RGBColor(148, 163, 184)
        
    # 1. HORAS 1
    b1 = plot_barra_global(["OFERTADAS", "ESTIMADAS", "REALES"], [metrics['h_ofe'], metrics['h_est'], metrics['h_real']], False, "Horas Hombre (HH)")
    agregar_diapositiva_grafica("HORAS GLOBALES DEL PROYECTO", "Comparativa general de Horas Hombre totales", b1, "HORAS 1")
    
    # 2. HORAS 2: Tarea
    b2 = plot_barras_agrupadas(df_f, 'Recurso', 'Horas', False)
    agregar_diapositiva_grafica("HORAS HOMBRE POR TAREA", "Distribución de HH por especialidad y puesto", b2, "HORAS 2")
    
    # 3. HORAS 2: GDF
    b3 = plot_barras_agrupadas(df_f, 'GDF', 'Horas', False)
    agregar_diapositiva_grafica("HORAS HOMBRE POR GDF", "Distribución de HH por subconjunto", b3, "HORAS 2")
    
    # 4. ESTIMADAS VS OFERTADAS: Cascada GDF
    b4 = plot_cascada(df_f, 'GDF', 'Horas', 'ESTIMADAS', 'OFERTADAS')
    agregar_diapositiva_grafica("HH ESTIMADO VS OFERTADO POR GDF", "Variación neta en horas por subconjunto", b4, "ESTIMADAS VS OFERTADAS")
    
    # 5. ESTIMADAS VS OFERTADAS: Cascada Tarea
    b5 = plot_cascada(df_f, 'Recurso', 'Horas', 'ESTIMADAS', 'OFERTADAS')
    agregar_diapositiva_grafica("HH ESTIMADO VS OFERTADO POR TAREA", "Variación neta en horas por disciplina", b5, "ESTIMADAS VS OFERTADAS")
    
    # 6. REALES VS ESTIMADAS: Cascada GDF
    b6 = plot_cascada(df_f, 'GDF', 'Horas', 'REALES', 'ESTIMADAS')
    agregar_diapositiva_grafica("HH REALES VS ESTIMADAS POR GDF", "Desviación en horas respecto a la estimación", b6, "REALES VS ESTIMADAS")
    
    # 7. REALES VS ESTIMADAS: Cascada Tarea
    b7 = plot_cascada(df_f, 'Recurso', 'Horas', 'REALES', 'ESTIMADAS')
    agregar_diapositiva_grafica("HH REALES VS ESTIMADAS POR TAREA", "Desviación en horas por disciplina", b7, "REALES VS ESTIMADAS")
    
    # 8. REALES VS OFERTADAS: Cascada GDF
    b8 = plot_cascada(df_f, 'GDF', 'Horas', 'REALES', 'OFERTADAS')
    agregar_diapositiva_grafica("HH REALES VS OFERTADAS POR GDF", "Balance final de horas por subconjunto", b8, "REALES VS OFERTADAS")
    
    # 9. REALES VS OFERTADAS: Cascada Tarea
    b9 = plot_cascada(df_f, 'Recurso', 'Horas', 'REALES', 'OFERTADAS')
    agregar_diapositiva_grafica("VARIACIÓN HH REALES VS OFERTADAS POR TAREA", "Balance final de horas por disciplina", b9, "REALES VS OFERTADAS")
    
    # 10. COSTOS: Global
    b10 = plot_barra_global(["OFERTADOS", "ESTIMADOS", "REALES"], [metrics['c_ofe'], metrics['c_est'], metrics['c_real']], True, "Costo Total ($)")
    agregar_diapositiva_grafica("COSTOS TOTALES DE FABRICACIÓN", "Impacto financiero global (Tarifas FABRICOST)", b10, "COSTOS")
    
    # 11. COSTOS 2: Recurso
    b11 = plot_barras_agrupadas(df_f, 'Recurso', 'Costo', True)
    agregar_diapositiva_grafica("COSTOS POR RECURSO / TAREA", "Distribución de costo por disciplina operativa", b11, "COSTOS 2")
    
    # 12. COSTOS 2: GDF
    b12 = plot_barras_agrupadas(df_f, 'GDF', 'Costo', True)
    agregar_diapositiva_grafica("COSTOS POR GDF (SUBCONJUNTO)", "Distribución de costo financiero por subconjunto", b12, "COSTOS 2")
    
    # 13. COSTOS ESTIM VS OFERT: Cascada GDF
    b13 = plot_cascada(df_f, 'GDF', 'Costo', 'ESTIMADAS', 'OFERTADAS')
    agregar_diapositiva_grafica("COSTOS ESTIMADOS VS OFERTADOS POR GDF", "Variación económica estimada por subconjunto", b13, "COSTOS ESTIM VS OFERT")
    
    # 14. COSTOS ESTIM VS OFERT: Cascada Tarea
    b14 = plot_cascada(df_f, 'Recurso', 'Costo', 'ESTIMADAS', 'OFERTADAS')
    agregar_diapositiva_grafica("COSTOS ESTIMADOS VS OFERTADOS POR TAREA", "Variación económica estimada por disciplina", b14, "COSTOS ESTIM VS OFERT")
    
    # 15. COSTOS REAL VS ESTIM: Cascada GDF
    b15 = plot_cascada(df_f, 'GDF', 'Costo', 'REALES', 'ESTIMADAS')
    agregar_diapositiva_grafica("COSTOS REALES VS ESTIMADOS POR GDF", "Desviación financiera real por subconjunto", b15, "COSTOS REAL VS ESTIM")
    
    # 16. COSTOS REAL VS ESTIM: Cascada Recurso
    b16 = plot_cascada(df_f, 'Recurso', 'Costo', 'REALES', 'ESTIMADAS')
    agregar_diapositiva_grafica("COSTOS REALES VS ESTIMADOS POR RECURSO", "Desviación financiera real por disciplina", b16, "COSTOS REAL VS ESTIM")
    
    # 17. COSTOS REAL VS OF: Cascada GDF
    b17 = plot_cascada(df_f, 'GDF', 'Costo', 'REALES', 'OFERTADAS')
    agregar_diapositiva_grafica("COSTOS REALES VS OFERTADOS POR GDF", "Balance económico final por subconjunto", b17, "COSTOS REAL VS OF")
    
    # 18. COSTOS REAL VS OF: Cascada Recurso
    b18 = plot_cascada(df_f, 'Recurso', 'Costo', 'REALES', 'OFERTADAS')
    agregar_diapositiva_grafica("COSTOS REALES VS OFERTADOS POR RECURSO", "Balance económico final por disciplina", b18, "COSTOS REAL VS OF")
    
    # 19. RENDIMIENTOS: Global
    b19 = plot_barra_global(["OFERTADO", "ESTIMADO", "REAL"], [metrics['rend_ofe'], metrics['rend_est'], metrics['rend_real']], False, "Kg / HH")
    agregar_diapositiva_grafica("RENDIMIENTOS OPERACIONALES (PESO / HH)", "Evolución del ratio de productividad por hora", b19, "RENDIMIENTOS")

    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    return out

if uploaded_file:
    file_bytes = io.BytesIO(uploaded_file.read())
    wb_check = openpyxl.load_workbook(file_bytes, read_only=True)
    hojas = wb_check.sheetnames
    wb_check.close()
    
    file_bytes.seek(0)
    df_tarifas = leer_fabricost(file_bytes)
    
    hojas_datos = [h for h in hojas if h != "FABRICOST"]
    hoja = st.sidebar.selectbox("Hoja de Proyecto activa:", hojas_datos, index=0)
    
    file_bytes.seek(0)
    df_raw = leer_matriz_exacta(file_bytes, hoja)
    df_long, df_pesos = procesar_datos(df_raw, df_tarifas)
    
    # --- FILTROS ---
    st.sidebar.markdown("---")
    st.sidebar.subheader("🔍 Filtros Dinámicos")
    
    lista_recursos = sorted(df_long['Recurso'].unique().tolist())
    filtro_recurso = st.sidebar.multiselect("RECURSO:", options=lista_recursos, default=[])
    
    lista_gdf = sorted(df_long['GDF'].dropna().unique().tolist())
    filtro_gdf = st.sidebar.multiselect("GDF:", options=lista_gdf, default=[])
    
    df_f = df_long.copy()
    df_p_f = df_pesos.copy()
    if filtro_recurso:
        df_f = df_f[df_f['Recurso'].isin(filtro_recurso)]
    if filtro_gdf:
        df_f = df_f[df_f['GDF'].isin(filtro_gdf)]
        df_p_f = df_p_f[df_p_f['GDF'].isin(filtro_gdf)]
        
    tot_horas = df_f.groupby('Estado')['Horas'].sum().to_dict()
    h_ofe = tot_horas.get("OFERTADAS", 0)
    h_est = tot_horas.get("ESTIMADAS", 0)
    h_real = tot_horas.get("REALES", 0)
    
    tot_costos = df_f.groupby('Estado')['Costo'].sum().to_dict()
    c_ofe = tot_costos.get("OFERTADAS", 0)
    c_est = tot_costos.get("ESTIMADAS", 0)
    c_real = tot_costos.get("REALES", 0)
    
    tot_peso = df_p_f.groupby('Estado')['Peso'].sum().to_dict()
    p_ofe = tot_peso.get("OFERTADAS", 0)
    p_est = tot_peso.get("ESTIMADAS", 0)
    p_real = tot_peso.get("REALES", 0)
    
    rend_ofe = (p_ofe / h_ofe) if h_ofe > 0 else 0
    rend_est = (p_est / h_est) if h_est > 0 else 0
    rend_real = (p_real / h_real) if h_real > 0 else 0
    
    v_eo = h_est - h_ofe
    p_eo = (v_eo / h_ofe * 100) if h_ofe else 0
    v_re = h_real - h_est
    p_re = (v_re / h_est * 100) if h_est else 0
    v_ro = h_real - h_ofe
    p_ro = (v_ro / h_ofe * 100) if h_ofe else 0
    
    vc_eo = c_est - c_ofe
    pc_eo = (vc_eo / c_ofe * 100) if c_ofe else 0
    vc_re = c_real - c_est
    pc_re = (vc_re / c_est * 100) if c_est else 0
    vc_ro = c_real - c_ofe
    pc_ro = (vc_ro / c_ofe * 100) if c_ofe else 0
    
    d_rend_eo = rend_est - rend_ofe
    p_rend_eo = (d_rend_eo / rend_ofe * 100) if rend_ofe else 0
    d_rend_re = rend_real - rend_est
    p_rend_re = (d_rend_re / rend_est * 100) if rend_est else 0
    d_rend_ro = rend_real - rend_ofe
    p_rend_ro = (d_rend_ro / rend_ofe * 100) if rend_ofe else 0

    dict_metrics = {
        'h_ofe': h_ofe, 'h_est': h_est, 'h_real': h_real,
        'v_eo': v_eo, 'p_eo': p_eo, 'v_re': v_re, 'p_re': p_re, 'v_ro': v_ro, 'p_ro': p_ro,
        'c_ofe': c_ofe, 'c_est': c_est, 'c_real': c_real,
        'vc_eo': vc_eo, 'pc_eo': pc_eo, 'vc_re': vc_re, 'pc_re': pc_re, 'vc_ro': vc_ro, 'pc_ro': pc_ro,
        'rend_ofe': rend_ofe, 'rend_est': rend_est, 'rend_real': rend_real,
        'd_rend_eo': d_rend_eo, 'p_rend_eo': p_rend_eo, 'd_rend_re': d_rend_re, 'p_rend_re': p_rend_re,
        'd_rend_ro': d_rend_ro, 'p_rend_ro': p_rend_ro
    }

    # --- BOTÓN DE EXPORTACIÓN PPTX (SIDEBAR) ---
    st.sidebar.markdown("---")
    st.sidebar.subheader("📥 Exportación a PowerPoint")
    
    pptx_bytes = generar_powerpoint_completo(df_f, dict_metrics, hoja)
    
    st.sidebar.download_button(
        label="📊 Descargar Presentación (.pptx)",
        data=pptx_bytes,
        file_name=f"Reporte_Completo_{hoja}.pptx",
        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )

    color_map = {"OFERTADAS": "#F59E0B", "ESTIMADAS": "#06B6D4", "REALES": "#6366F1"}
    orden_estados = ["OFERTADAS", "ESTIMADAS", "REALES"]

    def crear_cascada(df_base, col_agrupacion, metrica, estado_fin, estado_ini, titulo):
        piv = df_base.pivot_table(index=col_agrupacion, columns='Estado', values=metrica, aggfunc='sum').fillna(0)
        if estado_fin in piv.columns and estado_ini in piv.columns:
            piv['Desvio'] = piv[estado_fin] - piv[estado_ini]
            piv = piv.sort_values(by='Desvio', ascending=False).reset_index()
            
            y_vals = piv['Desvio'].tolist()
            y_total = sum(y_vals)
            text_labels = [formato_m_k(v) if metrica == 'Costo' else f"{int(v):,}" for v in y_vals] + [formato_m_k(y_total) if metrica == 'Costo' else f"{int(y_total):,}"]
            
            fig = go.Figure(go.Waterfall(
                orientation="v",
                measure=["relative"] * len(piv) + ["total"],
                x=piv[col_agrupacion].tolist() + ["Total"],
                y=y_vals + [y_total],
                text=text_labels,
                textposition="outside",
                decreasing={"marker": {"color": "#10B981"}},
                increasing={"marker": {"color": "#EF4444"}},
                totals={"marker": {"color": "#06B6D4"}}
            ))
            fig.update_layout(title=titulo, height=420, xaxis_tickangle=-45, yaxis_title="")
            return fig
        return None

    # --- LAS 11 PESTAÑAS COMPLETAS ---
    tabs = st.tabs([
        "HORAS 1", "HORAS 2", "ESTIMADAS VS OFERTADAS", "REALES VS ESTIMADAS", "REALES VS OFERTADAS",
        "COSTOS", "COSTOS 2", "COSTOS ESTIM VS OFERT", "COSTOS REAL VS ESTIM", "COSTOS REAL VS OF",
        "RENDIMIENTOS"
    ])
    
    with tabs[0]:
        st.subheader("HORAS")
        df_bh = pd.DataFrame({"Estado": orden_estados, "Horas": [h_ofe, h_est, h_real]})
        fig_h = px.bar(df_bh, x="Estado", y="Horas", color="Estado", text_auto=",.0f", color_discrete_map=color_map)
        fig_h.update_layout(height=400, showlegend=False, xaxis_title="", yaxis_title="Horas")
        st.plotly_chart(fig_h, use_container_width=True)
        
        c1, c2, c3 = st.columns(3)
        c1.metric("VARIACION HH ESTIMADO VS OFERTADO", f"{int(v_eo):,}", f"{p_eo:.2f}%")
        c2.metric("VARIACION HH REAL VS ESTIMADO", f"{int(v_re):,}", f"{p_re:.2f}%")
        c3.metric("VARIACION HH REALES VS OFERTADAS", f"{int(v_ro):,}", f"{p_ro:.2f}%")

    with tabs[1]:
        st.subheader("HORAS POR TAREA")
        df_t = df_f.groupby(['Recurso', 'Estado'])['Horas'].sum().reset_index()
        fig_t = px.bar(df_t, x="Recurso", y="Horas", color="Estado", barmode="group", category_orders={"Estado": orden_estados}, color_discrete_map=color_map)
        fig_t.update_layout(height=420, xaxis_tickangle=-45, yaxis_title="", legend_title="")
        st.plotly_chart(fig_t, use_container_width=True)
        
        st.subheader("HORAS POR GDF")
        df_g = df_f.groupby(['GDF', 'Estado'])['Horas'].sum().reset_index()
        fig_g = px.bar(df_g, x="GDF", y="Horas", color="Estado", barmode="group", category_orders={"Estado": orden_estados}, color_discrete_map=color_map)
        fig_g.update_layout(height=420, yaxis_title="", legend_title="")
        st.plotly_chart(fig_g, use_container_width=True)

    with tabs[2]:
        fig1 = crear_cascada(df_f, 'GDF', 'Horas', 'ESTIMADAS', 'OFERTADAS', "HH ESTIMADO VS OFERTADO POR GDF")
        if fig1: st.plotly_chart(fig1, use_container_width=True)
        fig2 = crear_cascada(df_f, 'Recurso', 'Horas', 'ESTIMADAS', 'OFERTADAS', "HH ESTIMADO VS OFERTADO POR TAREA")
        if fig2: st.plotly_chart(fig2, use_container_width=True)

    with tabs[3]:
        fig1 = crear_cascada(df_f, 'GDF', 'Horas', 'REALES', 'ESTIMADAS', "HH REALES VS ESTIMADAS POR GDF")
        if fig1: st.plotly_chart(fig1, use_container_width=True)
        fig2 = crear_cascada(df_f, 'Recurso', 'Horas', 'REALES', 'ESTIMADAS', "HH REALES VS ESTIMADAS POR TAREA")
        if fig2: st.plotly_chart(fig2, use_container_width=True)

    with tabs[4]:
        fig1 = crear_cascada(df_f, 'GDF', 'Horas', 'REALES', 'OFERTADAS', "HH REALES VS OFERTADAS POR GDF")
        if fig1: st.plotly_chart(fig1, use_container_width=True)
        fig2 = crear_cascada(df_f, 'Recurso', 'Horas', 'REALES', 'OFERTADAS', "VARIACION HH REALES VS OFERTADAS by RECURSO")
        if fig2: st.plotly_chart(fig2, use_container_width=True)

    with tabs[5]:
        st.subheader("COSTOS")
        df_bc = pd.DataFrame({"Estado": orden_estados, "Costo": [c_ofe, c_est, c_real]})
        df_bc["Etiqueta"] = df_bc["Costo"].apply(formato_m_k)
        fig_c = px.bar(df_bc, x="Estado", y="Costo", color="Estado", text="Etiqueta", color_discrete_map=color_map)
        fig_c.update_layout(height=400, showlegend=False, xaxis_title="", yaxis_title="Costo Total")
        st.plotly_chart(fig_c, use_container_width=True)
        
        c1, c2, c3 = st.columns(3)
        c1.metric("COSTOS ESTIMADOS VS OFERTADOS", formato_m_k(vc_eo), f"{pc_eo:.2f}%")
        c2.metric("COSTOS REALES VS ESTIMADOS", formato_m_k(vc_re), f"{pc_re:.2f}%")
        c3.metric("COSTOS REALES VS OFERTADOS", formato_m_k(vc_ro), f"{pc_ro:.2f}%")

    with tabs[6]:
        st.subheader("COSTOS POR RECURSO")
        df_cr = df_f.groupby(['Recurso', 'Estado'])['Costo'].sum().reset_index()
        fig_cr = px.bar(df_cr, x="Recurso", y="Costo", color="Estado", barmode="group", category_orders={"Estado": orden_estados}, color_discrete_map=color_map)
        fig_cr.update_layout(height=420, xaxis_tickangle=-45, yaxis_title="", legend_title="")
        st.plotly_chart(fig_cr, use_container_width=True)
        
        st.subheader("COSTOS POR GDF")
        df_cg = df_f.groupby(['GDF', 'Estado'])['Costo'].sum().reset_index()
        fig_cg = px.bar(df_cg, x="GDF", y="Costo", color="Estado", barmode="group", category_orders={"Estado": orden_estados}, color_discrete_map=color_map)
        fig_cg.update_layout(height=420, yaxis_title="", legend_title="")
        st.plotly_chart(fig_cg, use_container_width=True)

    with tabs[7]:
        fig1 = crear_cascada(df_f, 'GDF', 'Costo', 'ESTIMADAS', 'OFERTADAS', "COSTOS ESTIMADOS VS OFERTADOS POR GDF")
        if fig1: st.plotly_chart(fig1, use_container_width=True)
        fig2 = crear_cascada(df_f, 'Recurso', 'Costo', 'ESTIMADAS', 'OFERTADAS', "COSTOS ESTIMADOS VS OFERTADOS POR TAREA")
        if fig2: st.plotly_chart(fig2, use_container_width=True)

    with tabs[8]:
        fig1 = crear_cascada(df_f, 'GDF', 'Costo', 'REALES', 'ESTIMADAS', "COSTOS REALES VS ESTIMADOS POR GDF")
        if fig1: st.plotly_chart(fig1, use_container_width=True)
        fig2 = crear_cascada(df_f, 'Recurso', 'Costo', 'REALES', 'ESTIMADAS', "COSTOS REALES VS ESTIMADOS POR RECURSO")
        if fig2: st.plotly_chart(fig2, use_container_width=True)

    with tabs[9]:
        fig1 = crear_cascada(df_f, 'GDF', 'Costo', 'REALES', 'OFERTADAS', "COSTOS REALES VS OFERTADOS POR GDF")
        if fig1: st.plotly_chart(fig1, use_container_width=True)
        fig2 = crear_cascada(df_f, 'Recurso', 'Costo', 'REALES', 'OFERTADAS', "COSTOS REALES VS OFERTADOS POR RECURSO")
        if fig2: st.plotly_chart(fig2, use_container_width=True)

    with tabs[10]:
        st.subheader("RENDIMIENTOS (Peso / HH)")
        df_rend = pd.DataFrame({
            "Estado": ["OFERTADO", "ESTIMADO", "REAL"],
            "Rendimiento": [rend_ofe, rend_est, rend_real]
        })
        fig_r = px.bar(df_rend, x="Estado", y="Rendimiento", color="Estado", text=df_rend["Rendimiento"].apply(lambda v: f"{v:.2f}"), color_discrete_map={"OFERTADO": "#F59E0B", "ESTIMADO": "#06B6D4", "REAL": "#6366F1"})
        fig_r.update_layout(height=400, showlegend=False, xaxis_title="", yaxis_title="Kg / HH")
        st.plotly_chart(fig_r, use_container_width=True)
        
        cr1, cr2, cr3 = st.columns(3)
        cr1.metric("REND ESTIMADO VS OFERTADO", f"{d_rend_eo:.2f}", f"{p_rend_eo:.2f}% (%V E-O)")
        cr2.metric("REND REAL VS ESTIM", f"{d_rend_re:.2f}", f"{p_rend_re:.2f}% (%V R-E)")
        cr3.metric("REND REAL VS OFERTADO", f"{d_rend_ro:.2f}", f"{p_rend_ro:.2f}% (%V R-O)")
